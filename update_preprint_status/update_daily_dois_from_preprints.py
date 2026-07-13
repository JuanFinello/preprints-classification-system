#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Check preprint DOIs in daily_dois.xlsx against Crossref and optionally update
that same Excel file when a preprint has a published DOI.

The script does two things in one run:

1) Crossref check / report generation
   - Finds rows in daily_dois.xlsx where Status contains "preprint".
   - Queries Crossref for each unique DOI.
   - Generates CSV reports:
       preprints_checked_crossref.csv
       preprints_with_published_doi.csv
       significant_preprints_changed_to_published.csv

2) Optional Excel update
   - If --update-xlsx is used, updates daily_dois.xlsx in place.
   - Creates a backup before saving.
   - Replaces preprint DOI with published DOI.
   - If Observations is not empty:
       copies Observations -> Accessions
       sets Status = Reviewed
     Else:
       sets Status = none
   - Sets the current date in "From preprint to published ".

Example:
    python3 update_daily_dois_from_preprints.py \
        --input daily_dois.xlsx \
        --update-xlsx

Dry run example:
    python3 update_daily_dois_from_preprints.py \
        --input daily_dois.xlsx \
        --update-xlsx \
        --dry-run
"""

import argparse
import shutil
import time
from datetime import datetime
from pathlib import Path

import pandas as pd
import requests
from openpyxl import load_workbook


CROSSREF = "https://api.crossref.org/works/{}"
USER_AGENT = "JuanFinello-GISAID-preprint-check/1.0 (mailto:juanfinello@gmail.com)"


def norm_doi(x: str) -> str:
    """
    Normalize DOI strings for comparison:
    - strip
    - lower
    - remove https://doi.org/ or dx.doi.org prefixes
    - remove initial doi:
    """
    if x is None:
        return ""

    s = str(x).strip().lower()

    if not s:
        return ""

    for prefix in (
        "https://doi.org/",
        "http://doi.org/",
        "https://dx.doi.org/",
        "http://dx.doi.org/",
    ):
        if s.startswith(prefix):
            s = s[len(prefix):].strip()

    if s.startswith("doi:"):
        s = s[4:].strip()

    return s


def extract_related_dois(cr_item):
    """
    Extract related DOIs from Crossref, especially 'is-preprint-of'.
    """
    rel = (cr_item or {}).get("relation") or {}
    out = {"is_preprint_of": []}

    vals = rel.get("is-preprint-of") or []

    for v in vals:
        if isinstance(v, dict):
            rid = v.get("id") or v.get("DOI") or v.get("doi")
            if rid:
                out["is_preprint_of"].append(norm_doi(rid))
        elif isinstance(v, str):
            out["is_preprint_of"].append(norm_doi(v))

    # Unique values, preserving order
    seen = set()
    out["is_preprint_of"] = [
        x for x in out["is_preprint_of"]
        if x and not (x in seen or seen.add(x))
    ]

    return out


def classify_doi(doi, session):
    """
    Classify a DOI using Crossref.

    Returns one of:
        published
        preprint
        unknown
    """
    doi_norm = norm_doi(doi)
    url = CROSSREF.format(requests.utils.quote(doi_norm, safe=""))

    try:
        r = session.get(url, timeout=30)

        if r.status_code == 404:
            return {
                "doi": doi_norm,
                "state": "unknown",
                "published_doi": "",
                "reason": "404",
            }

        r.raise_for_status()
        item = r.json().get("message", {})

    except Exception as e:
        return {
            "doi": doi_norm,
            "state": "unknown",
            "published_doi": "",
            "reason": type(e).__name__,
        }

    cr_type = (item.get("type") or "").lower()
    related = extract_related_dois(item)

    # Explicit relation: preprint -> published article
    if related["is_preprint_of"]:
        return {
            "doi": doi_norm,
            "state": "published",
            "published_doi": related["is_preprint_of"][0],
            "reason": "relation_is_preprint_of",
        }

    # Already a published article
    if cr_type in {"journal-article", "proceedings-article"}:
        return {
            "doi": doi_norm,
            "state": "published",
            "published_doi": doi_norm,
            "reason": "crossref_type_published",
        }

    # Still a preprint
    if cr_type in {"posted-content", "preprint"}:
        return {
            "doi": doi_norm,
            "state": "preprint",
            "published_doi": "",
            "reason": "crossref_type_preprint",
        }

    return {
        "doi": doi_norm,
        "state": "unknown",
        "published_doi": "",
        "reason": "no_signal",
    }


def find_header_map(header_row) -> dict[str, int]:
    """
    Return exact column name -> 1-based openpyxl column index.
    """
    col_map = {}

    for cell in header_row:
        if cell.value is None:
            continue
        col_map[str(cell.value)] = cell.column

    return col_map


def update_daily_dois_xlsx(
    xlsx_path: Path,
    df_pub: pd.DataFrame,
    sheet: str | None = None,
    dry_run: bool = False,
) -> None:
    """
    Update daily_dois.xlsx using the published DOI mappings detected by Crossref.
    """
    if df_pub.empty:
        print("No published DOI mappings detected. XLSX update skipped.")
        return

    required_pub_cols = {"DOI", "published_doi"}
    missing_pub = required_pub_cols - set(df_pub.columns)

    if missing_pub:
        raise ValueError(
            f"Published DOI dataframe is missing columns: {missing_pub}. "
            f"Available columns: {list(df_pub.columns)}"
        )

    pre_to_pub = {}

    for _, row in df_pub.iterrows():
        pre = norm_doi(row.get("DOI", ""))
        pub = norm_doi(row.get("published_doi", ""))

        if pre and pub:
            pre_to_pub[pre] = pub

    print(f"Mappings preprint->published available for XLSX update: {len(pre_to_pub)}")

    if not pre_to_pub:
        print("No valid preprint->published mappings. XLSX update skipped.")
        return

    backup_path = xlsx_path.with_suffix(xlsx_path.suffix + ".bak")

    if not dry_run:
        shutil.copy2(xlsx_path, backup_path)
        print(f"Backup created: {backup_path}")

    wb = load_workbook(xlsx_path)
    ws = wb[sheet] if sheet else wb.active

    header_row = ws[1]
    col_map = find_header_map(header_row)

    required_cols = [
        "DOI",
        "Observations",
        "Accessions",
        "Status",
        "From preprint to published ",
    ]

    missing = [c for c in required_cols if c not in col_map]

    if missing:
        raise ValueError(
            "Missing required columns in XLSX header row 1: "
            + ", ".join(missing)
            + f"\nAvailable columns: {list(col_map.keys())}"
        )

    col_doi = col_map["DOI"]
    col_obs = col_map["Observations"]
    col_acc = col_map["Accessions"]
    col_status = col_map["Status"]
    col_stamp = col_map["From preprint to published "]

    today = datetime.now().strftime("%d/%m/%Y")

    matched = 0
    updated_reviewed = 0
    updated_none = 0
    copied_obs = 0
    doi_replaced = 0

    for r in range(2, ws.max_row + 1):
        doi_cell_val = ws.cell(row=r, column=col_doi).value
        doi_norm = norm_doi(doi_cell_val)

        if not doi_norm:
            continue

        if doi_norm not in pre_to_pub:
            continue

        matched += 1
        published_norm = pre_to_pub[doi_norm]

        # Replace DOI if published_doi is different
        if published_norm and published_norm != doi_norm:
            ws.cell(row=r, column=col_doi).value = published_norm
            doi_replaced += 1

        obs_val = ws.cell(row=r, column=col_obs).value
        obs_s = "" if obs_val is None else str(obs_val).strip()

        if obs_s != "":
            # Copy Observations -> Accessions, without deleting Observations
            ws.cell(row=r, column=col_acc).value = obs_val
            ws.cell(row=r, column=col_status).value = "Reviewed"
            updated_reviewed += 1
            copied_obs += 1
        else:
            ws.cell(row=r, column=col_status).value = "none"
            updated_none += 1

        # Always set timestamp
        ws.cell(row=r, column=col_stamp).value = today

    print("-" * 60)
    print(f"Rows with matching preprint DOI in XLSX: {matched}")
    print(f"DOIs replaced by published DOI: {doi_replaced}")
    print(f"Status='Reviewed' because Observations was not empty: {updated_reviewed}")
    print(f"Status='none' because Observations was empty: {updated_none}")
    print(f"Observations copied to Accessions: {copied_obs}")
    print(f"Timestamp set in 'From preprint to published ': {today}")
    print("-" * 60)

    if dry_run:
        print("Dry-run: no XLSX changes were saved.")
        return

    wb.save(xlsx_path)
    print(f"XLSX updated: {xlsx_path}")


def main():
    parser = argparse.ArgumentParser(
        description=(
            "Check preprint DOIs from daily_dois.xlsx against Crossref, "
            "generate CSV reports, and optionally update the XLSX."
        )
    )

    parser.add_argument(
        "-i",
        "--input",
        default="daily_dois.xlsx",
        help="Input Excel file. Default: daily_dois.xlsx",
    )

    parser.add_argument(
        "--sheet",
        default=None,
        help="Sheet name to read/update. If omitted, uses the first/active sheet.",
    )

    parser.add_argument(
        "--all-output",
        default="preprints_checked_crossref.csv",
        help="CSV output with all checked preprints.",
    )

    parser.add_argument(
        "--published-output",
        default="preprints_with_published_doi.csv",
        help="CSV output only with preprints that have a published DOI.",
    )

    parser.add_argument(
        "--significant-output",
        default="significant_preprints_changed_to_published.csv",
        help="CSV output only with significant preprints that now have a published DOI.",
    )

    parser.add_argument(
        "--sleep",
        type=float,
        default=0.2,
        help="Seconds between Crossref requests. Default: 0.2",
    )

    parser.add_argument(
        "--update-xlsx",
        action="store_true",
        help="Update the input XLSX in place using the detected published DOI mappings.",
    )

    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Run the full process but do not save XLSX updates. CSV reports are still written.",
    )

    args = parser.parse_args()

    daily_dois_file = Path(args.input)

    if not daily_dois_file.exists():
        raise FileNotFoundError(f"Input file not found: {daily_dois_file}")

    # =========================
    # Read daily_dois
    # =========================
    df_daily = pd.read_excel(daily_dois_file, sheet_name=args.sheet or 0, dtype=str)

    # =========================
    # Validate columns needed for Crossref check
    # =========================
    required_cols = {
        "DOI",
        "Status",
        "Significant_preprint",
    }

    missing = required_cols - set(df_daily.columns)

    if missing:
        raise ValueError(
            f"Missing columns: {missing}\n"
            f"Available columns: {list(df_daily.columns)}"
        )

    # =========================
    # Normalize key columns
    # =========================
    df_daily["DOI"] = df_daily["DOI"].apply(norm_doi)

    df_daily["Status"] = (
        df_daily["Status"]
        .astype(str)
        .str.strip()
        .str.lower()
    )

    df_daily["Significant_preprint"] = (
        df_daily["Significant_preprint"]
        .astype(str)
        .str.strip()
        .str.lower()
    )

    # =========================
    # Filter preprints
    # =========================
    df_preprints = df_daily[
        df_daily["Status"].str.contains("preprint", na=False)
    ].copy()

    print(f"Preprints detected: {len(df_preprints)}")

    if df_preprints.empty:
        print("No preprints detected. No CSV reports generated and no XLSX update done.")
        return

    df_preprints = df_preprints[
        [
            "DOI",
            "Status",
            "Significant_preprint",
        ]
    ].copy()

    # =========================
    # Unique DOIs
    # =========================
    dois = (
        df_preprints["DOI"]
        .dropna()
        .astype(str)
        .str.strip()
        .loc[lambda s: s.ne("") & s.ne("nan")]
        .drop_duplicates()
        .tolist()
    )

    print(f"Unique DOIs to query in Crossref: {len(dois)}")

    # =========================
    # Crossref session
    # =========================
    session = requests.Session()
    session.headers.update({"User-Agent": USER_AGENT})

    # =========================
    # Query Crossref
    # =========================
    results = []

    for i, doi in enumerate(dois, 1):
        results.append(classify_doi(doi, session))

        if i % 25 == 0 or i == len(dois):
            print(f"Queried: {i}/{len(dois)}")

        time.sleep(args.sleep)

    res_df = pd.DataFrame(results)

    print("\nCrossref state summary:")
    print(res_df["state"].value_counts(dropna=False))

    # =========================
    # Merge results
    # =========================
    df_preprints["_doi_norm"] = df_preprints["DOI"].apply(norm_doi)
    res_df["_doi_norm"] = res_df["doi"].apply(norm_doi)

    df_preprints_checked = df_preprints.merge(
        res_df.drop(columns=["doi"]),
        on="_doi_norm",
        how="left",
    ).drop(columns=["_doi_norm"])

    # =========================
    # Preprints with published DOI
    # =========================
    df_pub = df_preprints_checked[
        df_preprints_checked["published_doi"].notna()
        & (
            df_preprints_checked["published_doi"]
            .astype(str)
            .str.strip()
            != ""
        )
    ].copy()

    # =========================
    # Significant preprints now published
    # =========================
    df_sig_pub = df_preprints_checked[
        (
            df_preprints_checked["Significant_preprint"]
            .astype(str)
            .str.strip()
            .str.lower()
            == "x"
        )
        & df_preprints_checked["published_doi"].notna()
        & (
            df_preprints_checked["published_doi"]
            .astype(str)
            .str.strip()
            != ""
        )
    ].copy()

    # =========================
    # Save CSV outputs
    # =========================
    df_preprints_checked.to_csv(
        args.all_output,
        sep=",",
        index=False,
        encoding="utf-8",
    )

    df_pub.to_csv(
        args.published_output,
        sep=",",
        index=False,
        encoding="utf-8",
    )

    df_sig_pub.to_csv(
        args.significant_output,
        sep=",",
        index=False,
        encoding="utf-8",
    )

    print(f"\nCSV with all checked preprints: {args.all_output}")
    print(f"CSV with preprints that have published DOI: {args.published_output}")
    print(f"CSV with significant preprints now published: {args.significant_output}")
    print(f"\nPublished preprints detected: {len(df_pub)}")
    print(f"Significant published preprints detected: {len(df_sig_pub)}")

    # =========================
    # Optional XLSX update
    # =========================
    if args.update_xlsx:
        print("\nUpdating XLSX using detected published DOI mappings...")
        update_daily_dois_xlsx(
            xlsx_path=daily_dois_file,
            df_pub=df_pub,
            sheet=args.sheet,
            dry_run=args.dry_run,
        )
    else:
        print("\nXLSX was not updated. Use --update-xlsx to update daily_dois.xlsx.")


if __name__ == "__main__":
    main()
